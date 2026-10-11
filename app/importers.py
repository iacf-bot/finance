"""CSV (bank) import and Suishouji (随手记) export import."""
from __future__ import annotations

import csv
import io
import re
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from typing import Any

from . import ledger
from .db import Database
from .parser import categorize

# ====================================================================== shared helpers
DATE_FORMATS = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d",
                "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y", "%m/%d/%y", "%d.%m.%Y", "%Y.%m.%d", "%Y%m%d",
                "%Y年%m月%d日 %H:%M:%S", "%Y年%m月%d日 %H:%M", "%Y年%m月%d日"]


def parse_date_any(v: Any, date_format: str | None = None) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, (int, float)):
        if 20000 < v < 80000:  # excel serial date
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).strftime("%Y-%m-%d %H:%M")
        v = str(int(v))
    s = str(v).strip().replace("\t", "")
    fmts = [date_format] if date_format else DATE_FORMATS
    for f in fmts:
        try:
            d = datetime.strptime(s, f)
            return d.strftime("%Y-%m-%d %H:%M") if ("%H" in f) else d.strftime("%Y-%m-%d 12:00")
        except ValueError:
            continue
    return None


def parse_money(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("¥", "").replace("￥", "").replace("$", "").replace("元", "").replace(" ", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if s.endswith("-"):
        neg, s = True, s[:-1]
    try:
        f = float(s)
    except ValueError:
        return None
    return -f if neg else f


def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ====================================================================== bank CSV import
FIELD_SYNONYMS = {
    "date": ["transaction date", "trans. date", "trans date", "date", "posted date", "post date", "posting date",
             "交易时间", "交易日期", "日期", "记账日期", "交易创建时间", "付款时间"],
    "amount": ["amount", "amount (usd)", "transaction amount", "金额", "金额(元)", "交易金额", "金额（元）"],
    "debit": ["debit", "withdrawal", "withdrawals", "withdrawal amount", "支出", "支出金额", "借方金额"],
    "credit": ["credit", "deposit", "deposits", "deposit amount", "收入", "存入金额", "收入金额", "贷方金额"],
    "description": ["description", "payee", "merchant", "name", "交易对方", "商户名称", "对方户名", "摘要", "交易摘要", "payee name"],
    "category": ["category", "交易分类", "分类"],
    "note": ["memo", "notes", "note", "备注", "商品说明", "商品", "extended details"],
    "direction": ["收/支", "收支", "type", "transaction type", "资金方向"],
    "currency": ["currency", "币种"],
}


def read_csv_table(data: bytes) -> tuple[list[str], list[list[str]]]:
    text = decode_bytes(data)
    lines = text.splitlines()
    # skip preamble (Alipay/WeChat exports start with several info lines)
    start = 0
    known = {s for v in FIELD_SYNONYMS.values() for s in v}
    for i, line in enumerate(lines[:60]):
        cells = [c.strip().strip('"').lower() for c in re.split(r"[,\t;]", line)]
        if sum(1 for c in cells if c in known) >= 2:
            start = i
            break
    body = "\n".join(lines[start:])
    try:
        dialect = csv.Sniffer().sniff(body[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(body), dialect)]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], []
    header = rows[0]
    data_rows = [r for r in rows[1:] if len([c for c in r if c]) >= 2 and not r[0].startswith("---")]
    return header, data_rows


def guess_mapping(header: list[str]) -> tuple[dict[str, str | None], str, str]:
    """Return (field->column name, detected format name, sign convention)."""
    low = {h.lower().strip(): h for h in header}
    mapping: dict[str, str | None] = {}
    for field, syns in FIELD_SYNONYMS.items():
        mapping[field] = next((low[s] for s in syns if s in low), None)
    hs = set(low)
    fmt, sign = "Generic", "negative_is_expense"
    if {"transaction date", "post date", "description", "category", "type", "amount"} <= hs:
        fmt = "Chase credit card"
    elif {"details", "posting date", "description", "amount"} <= hs:
        fmt = "Chase checking"
        mapping["date"] = low["posting date"]; mapping["direction"] = None
    elif {"date", "description", "amount", "running bal."} <= hs:
        fmt = "Bank of America checking"
    elif {"posted date", "payee", "address", "amount"} <= hs:
        fmt = "Bank of America credit card"
    elif {"transaction date", "posted date", "card no.", "description", "category", "debit", "credit"} <= hs:
        fmt = "Capital One"
    elif {"status", "date", "description", "debit", "credit"} <= hs:
        fmt = "Citi"
    elif {"trans. date", "post date", "description", "amount", "category"} <= hs:
        fmt, sign = "Discover", "positive_is_expense"
    elif {"date", "description", "amount"} <= hs and ("extended details" in hs or "card member" in hs or len(hs) <= 5):
        fmt, sign = ("American Express" if ("card member" in hs or "extended details" in hs) else "Generic (Date, Description, Amount)"), \
            ("positive_is_expense" if ("card member" in hs or "extended details" in hs) else "negative_is_expense")
    elif "收/支" in hs and ("交易对方" in hs or "商品说明" in hs or "商品" in hs):
        fmt, sign = ("WeChat Pay 微信支付" if "金额(元)" in hs else "Alipay 支付宝"), "direction_column"
        mapping["description"] = low.get("交易对方")
    if fmt == "Chase credit card":
        mapping["direction"] = None
    return mapping, fmt, sign


def csv_preview(data: bytes) -> dict:
    header, rows = read_csv_table(data)
    mapping, fmt, sign = guess_mapping(header)
    return {"header": header, "rows": rows[:10], "row_count": len(rows), "mapping": mapping, "format": fmt, "sign": sign}


def csv_commit(db: Database, data: bytes, account_id: int, mapping: dict, sign: str = "negative_is_expense",
               date_format: str | None = None, status: str = "posted", skip_duplicates: bool = True,
               ledger_id: int | None = None) -> dict:
    header, rows = read_csv_table(data)
    idx = {h: i for i, h in enumerate(header)}
    acct = db.one("SELECT * FROM accounts WHERE id=?", (account_id,))
    if not acct:
        raise ledger.ValidationError("Choose an account to import into")
    cats = db.all("SELECT * FROM categories")
    summary = {"imported": 0, "skipped": [], "duplicates": 0, "by_kind": Counter(), "review": 0}
    occ: Counter = Counter()

    def col(r, field):
        name = mapping.get(field)
        if not name or name not in idx or idx[name] >= len(r):
            return ""
        return r[idx[name]]

    for n, r in enumerate(rows, start=2):
        d = parse_date_any(col(r, "date"), date_format)
        if not d:
            summary["skipped"].append({"row": n, "reason": f"bad date '{col(r, 'date')}'"}); continue
        amt = parse_money(col(r, "amount"))
        if amt is None and (mapping.get("debit") or mapping.get("credit")):
            deb, cre = parse_money(col(r, "debit")) or 0, parse_money(col(r, "credit")) or 0
            amt = (cre - abs(deb)) if (deb or cre) else None
            row_sign = "negative_is_expense"
        else:
            row_sign = sign
        if amt is None or amt == 0:
            summary["skipped"].append({"row": n, "reason": "no amount"}); continue
        desc = col(r, "description")
        note = col(r, "note")
        direction = col(r, "direction")
        if row_sign == "direction_column":
            if "不计" in direction or direction.strip() in ("/", ""):
                summary["skipped"].append({"row": n, "reason": f"neutral row ({direction or 'no 收/支'})"}); continue
            inflow = "收" in direction or direction.lower() in ("credit", "income")
            amt = abs(amt) if inflow else -abs(amt)
        elif row_sign == "positive_is_expense":
            amt = -amt
        text = f"{desc} {note} {col(r, 'category')}"
        is_payment = bool(re.search(r"payment|autopay|thank you|还款", text, re.I)) and amt > 0 and acct["type"] == "credit_card"
        if is_payment:
            kind = "transfer"
        elif amt < 0:
            kind = "expense"
        else:
            kind = "refund" if re.search(r"refund|return|退款", text, re.I) else "income"
        cat_id = None
        if kind in ("expense", "refund", "income"):
            cat_id, _, _ = categorize(text, cats, kind="income" if kind == "income" else "expense")
        base_key = ledger.tx_dedupe_key(d, amt, str(account_id), desc + note)
        occ[base_key] += 1
        ext = f"csv:{base_key}#{occ[base_key]}"
        if skip_duplicates and db.one("SELECT 1 x FROM transactions WHERE external_id=?", (ext,)):
            summary["duplicates"] += 1; continue
        st = status
        tx = {"kind": kind, "date": d, "amount": abs(amt), "account_id": account_id, "category_id": cat_id,
              "merchant": desc[:80] or None, "note": note[:500], "source": "import", "status": st,
              "external_id": ext, "raw_text": ",".join(r)[:1000], "ledger_id": ledger_id or 1}
        if kind == "transfer":  # card payment: source account unknown -> review
            tx.update({"to_account_id": account_id, "account_id": None, "status": "review"})
        try:
            ledger.create_tx(db, tx)
        except ledger.ValidationError as e:
            tx["status"] = "review"
            try:
                ledger.create_tx(db, tx)
            except ledger.ValidationError:
                summary["skipped"].append({"row": n, "reason": str(e)}); continue
        summary["imported"] += 1
        summary["by_kind"][kind] += 1
        if tx["status"] == "review":
            summary["review"] += 1
    summary["by_kind"] = dict(summary["by_kind"])
    summary["skipped"] = summary["skipped"][:200]
    return summary


def export_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    cols = ["id", "date", "kind", "amount", "currency", "account_name", "to_account_name", "to_amount",
            "parent_category_name", "category_name", "merchant_name", "member_name", "project_name", "tags", "note",
            "source", "status", "ticker", "shares", "price", "fee", "orig_amount", "orig_currency", "ledger_name"]
    w = csv.writer(buf)
    w.writerow(cols)
    for r in rows:
        w.writerow(["" if r.get(c) is None else r.get(c) for c in cols])
    return "\ufeff" + buf.getvalue()


# ====================================================================== Suishouji import
# Column names confirmed from ezbookkeeping's Feidee importers (see docs/sui_reference.md).
SUI_COLUMNS = {
    "type": ["交易类型"],
    "date": ["日期", "交易时间", "时间"],
    "category": ["分类", "一级分类", "类别"],
    "subcategory": ["子分类", "二级分类", "子类别"],
    "account": ["账户1", "账户", "支出账户", "收入账户", "转出账户"],
    "account2": ["账户2", "转入账户"],
    "currency": ["账户币种", "币种"],
    "amount": ["金额"],
    "member": ["成员"],
    "merchant": ["商家"],
    "project": ["项目"],
    "project_category": ["项目分类"],
    "recorder": ["记账人"],
    "note": ["备注"],
    "related_id": ["关联Id", "关联ID", "关联id"],
}
REQUIRED = ["date", "account", "amount"]
SUI_TYPES = {
    "支出": "expense", "收入": "income", "转账": "transfer", "转出": "transfer_out", "转入": "transfer_in",
    "余额变更": "balance", "负债变更": "liability_change", "债务变更": "liability_change", "债权变更": "receivable_change",
    "退款": "refund", "借入": "borrow", "借出": "lend", "收债": "collect", "还债": "repay", "代付": "pay_for",
    "报销": "reimburse",
}
# For loan-type rows: (guessed type of account1, guessed type of account2). Directions are best effort.
LOAN_ROLES = {"borrow": ("payable", None), "repay": (None, "payable"), "lend": (None, "receivable"),
              "collect": ("receivable", None), "pay_for": (None, "receivable"), "reimburse": ("receivable", None)}


def guess_account_type(name: str, hint: str | None = None) -> str:
    """Guess an account type from its name. Latin keywords match on word boundaries (so "aspirate" is not "IRA")."""
    n = name or ""
    n = re.sub(r"(?<=[a-z])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", " ", n)  # XYChecking -> XY Checking, FakeRedCard -> Fake Red Card
    n = re.sub(r"(?<=[A-Za-z])(?=\d)|(?<=\d)(?=[A-Za-z]{2})", " ", n).lower()  # Fake2Save -> Fake 2 Save (keeps 401k)
    W = lambda words: r"(?<![a-z])(?:" + words + r")(?![a-z])"
    rules = [
        (W(r"debit"), "checking"),
        (r"信用卡|白条|花呗|贷记|" + W(r"credit|visa|master(?:card)?|amex|discover|redcard|card|freedom|sapphire|reward|rewards"), "credit_card"),
        (r"房贷|车贷|贷款|借呗|微粒贷|" + W(r"loan|mortgage|mortage"), "loan"),
        (r"应付|借入|欠款|" + W(r"payable"), "payable"),
        (r"应收|借出|报销|债权|垫付|" + W(r"receivable"), "receivable"),
        (r"养老|公积金|" + W(r"401k?|403b|ira|roth|pension|529|hsa inv"), "retirement"),
        (r"股票|证券|基金|理财|券商|期货|" + W(r"stock|stocks|broker|brokerage|fund|investment|invest|fidelity|vanguard|schwab|webull|robinhood|etrade|e\*trade|ibkr|merrill"), "brokerage"),
        (r"比特币|币安|" + W(r"btc|eth|crypto|coinbase"), "crypto"),
        (r"支付宝|微信|零钱|余额宝|钱包|饭卡|公交卡|交通卡|" + W(r"alipay|wechat|paypal|venmo|zelle|wallet"), "ewallet"),
        (r"现金|" + W(r"cash|usd|cny|rmb|jpy|eur|gbp|hkd|cad|twd|krw"), "cash"),
        (r"定期|储蓄账户|存单|" + W(r"savings?|save|cd"), "savings"),
        (r"银行|储蓄卡|借记|工行|招行|建行|农行|中行|交行|" + W(r"checking|bank|debit|wf|wells|boa|bofa|chase|citi|cmb|icbc|abc|ccb|boc"), "checking"),
    ]
    for rx, t in rules:
        if re.search(rx, n):
            return t
    return hint or "other_asset"


def guess_account_type_from_usage(name: str, stats: Counter, hint: str | None = None) -> str:
    """Use what the export says about the account (负债变更/债权变更 rows, spending) on top of the name."""
    by_name = guess_account_type(name)
    liability_rows = stats.get("liability_change", 0)
    if liability_rows:  # 随手记 only writes 负债变更 for liability accounts (credit cards, loans, payables)
        if by_name in ("credit_card", "loan", "payable"):
            return by_name
        return "credit_card" if stats.get("expense", 0) else "payable"
    if stats.get("receivable_change", 0):
        return "receivable"
    if stats.get("balance", 0) and by_name in ("credit_card", "loan", "payable"):
        return "other_asset"  # 随手记 writes 余额变更 only for asset accounts (e.g. a rewards/points or gift-card account)
    if hint in ("payable", "receivable") and by_name in ("other_asset", "checking"):
        return hint
    if by_name == "other_asset" and hint:
        return hint
    return by_name


def _read_sheets(filename: str, data: bytes) -> list[tuple[str, list[list[Any]]]]:
    name = (filename or "").lower()
    if data[:4] == b"PK\x03\x04" or name.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or name.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        out = []
        for sh in wb.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, wb.datemode))
                    else:
                        row.append(cell.value)
                rows.append(row)
            out.append((sh.name, rows))
        return out
    # CSV (app export: first line "随手记导出文件(headers:v5;...)")
    text = decode_bytes(data).lstrip("\ufeff")
    if text.startswith("随手记导出文件"):  # "随手记导出文件(headers:v5; iPhone; MyMoney; ...)" banner line
        text = text.split("\n", 1)[1] if "\n" in text else ""
    sample = text[:4096]
    delim = "\t" if sample.count("\t") > sample.count(",") else ","
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    return [("CSV", rows)]


def _detect_columns(header: list[Any]) -> tuple[dict[str, int], list[str]]:
    h = [str(x).strip() if x is not None else "" for x in header]
    cols: dict[str, int] = {}
    for field, syns in SUI_COLUMNS.items():
        for s in syns:
            if s in h:
                cols[field] = h.index(s)
                break
    used = set(cols.values())
    unknown = [x for i, x in enumerate(h) if x and i not in used]
    return cols, unknown


def _cell(row, i):
    if i is None or i >= len(row):
        return None
    v = row[i]
    if isinstance(v, str):
        v = v.strip()
    return v if v not in ("", None) else None


def parse_sui(filename: str, data: bytes, column_overrides: dict | None = None) -> dict:
    """Parse a Suishouji export into normalized rows. No DB access."""
    sheets_out, records, warnings = [], [], []
    for sname, rows in _read_sheets(filename, data):
        hdr_idx = next((i for i, r in enumerate(rows[:10]) if any(str(c or "").strip() in ("交易类型", "日期") for c in r)), None)
        if hdr_idx is None:
            sheets_out.append({"name": sname, "rows": 0, "skipped": True, "reason": "no 交易类型/日期 header", "header": []})
            continue
        header = [str(c).strip() if c is not None else "" for c in rows[hdr_idx]]
        cols, unknown = _detect_columns(header)
        for field, colname in ((column_overrides or {}).get(sname) or {}).items():
            if colname in header:
                cols[field] = header.index(colname)
        missing = [f for f in REQUIRED if f not in cols]
        info = {"name": sname, "header": header, "columns": {f: header[i] for f, i in cols.items()},
                "unrecognized": unknown, "missing": missing, "rows": 0, "types": Counter()}
        sheets_out.append(info)
        if missing:
            info["skipped"] = True
            info["reason"] = "missing columns: " + ", ".join(missing)
            continue
        for rn, r in enumerate(rows[hdr_idx + 1:], start=hdr_idx + 2):
            if not any(c not in (None, "") for c in r):
                continue
            t_raw = str(_cell(r, cols.get("type")) or sname).strip()
            info["rows"] += 1
            info["types"][t_raw] += 1
            records.append({"sheet": sname, "row": rn, "type_raw": t_raw, "type": SUI_TYPES.get(t_raw),
                            "date_raw": _cell(r, cols.get("date")), "date": parse_date_any(_cell(r, cols.get("date"))),
                            "category": _cell(r, cols.get("category")), "subcategory": _cell(r, cols.get("subcategory")),
                            "account": str(_cell(r, cols.get("account")) or "") or None,
                            "account2": str(_cell(r, cols.get("account2")) or "") or None,
                            "currency": _cell(r, cols.get("currency")), "amount": parse_money(_cell(r, cols.get("amount"))),
                            "member": _cell(r, cols.get("member")), "merchant": _cell(r, cols.get("merchant")),
                            "project": _cell(r, cols.get("project")), "note": _cell(r, cols.get("note")),
                            "related_id": _cell(r, cols.get("related_id"))})
        info["types"] = dict(info["types"])
    # pair app-CSV 转出/转入 rows by 关联Id
    pending: dict[str, dict] = {}
    merged = []
    for rec in records:
        if rec["type"] in ("transfer_out", "transfer_in"):
            rid = rec.get("related_id")
            if not rid:
                rec["error"] = "transfer half without 关联Id"; merged.append(rec); continue
            other = pending.pop(rid, None)
            if not other:
                pending[rid] = rec; continue
            out_r, in_r = (rec, other) if rec["type"] == "transfer_out" else (other, rec)
            m = dict(out_r); m["type"] = "transfer"; m["type_raw"] = "转账"; m["account2"] = in_r["account"]
            m["row"] = f"{out_r['row']}+{in_r['row']}"
            m["to_amount"] = in_r["amount"]; m["currency2"] = in_r.get("currency")
            merged.append(m)
        else:
            merged.append(rec)
    for rec in pending.values():
        rec["error"] = "transfer half without its matching row (关联Id)"
        merged.append(rec)
    return {"sheets": sheets_out, "records": merged, "warnings": warnings}



def ledger_is_liability(acct_type: str) -> bool:
    from .db import ACCOUNT_TYPES
    return bool(ACCOUNT_TYPES.get(acct_type, ("", "", 0, 0))[2])


def _record_effects(r: dict) -> list[tuple[str, float]]:
    """Signed effect of one parsed 随手记 record on account balances (asset-positive convention), for previews."""
    t, amt = r.get("type"), r.get("amount")
    if r.get("error") or not t or not amt or not r.get("account") or not r.get("date"):
        return []
    a1, a2 = r["account"], r.get("account2")
    if t == "expense":
        return [(a1, -amt)]
    if t in ("income", "refund"):
        return [(a1, amt if t == "income" else abs(amt))]
    if t == "balance" or t == "receivable_change":
        return [(a1, amt)]
    if t == "liability_change":
        return [(a1, -amt)]
    if (t == "transfer" or t in LOAN_ROLES) and a2 and a2 != a1:
        to_amt = r.get("to_amount")
        return [(a1, -abs(amt)), (a2, abs(to_amt) if to_amt is not None else abs(amt))]
    return []


def _account_usage(recs: list[dict]) -> dict[str, dict]:
    info: dict[str, dict] = {}
    for r in recs:
        roles = LOAN_ROLES.get(r.get("type") or "", (None, None))
        for key, hint in (("account", roles[0]), ("account2", roles[1])):
            nm = r.get(key)
            if not nm:
                continue
            a = info.setdefault(nm, {"name": nm, "currency": None, "rows": 0, "hint": None, "stats": Counter(),
                                     "net": 0.0, "first": None, "last": None})
            a["rows"] += 1
            a["stats"][(r.get("type") or "unknown") + ("" if key == "account" else "_in")] += 1
            cur = r.get("currency") if key == "account" else r.get("currency2")
            if cur:
                a["currency"] = str(cur).upper()
            if hint:
                a["hint"] = hint
            d = (r.get("date") or "")[:10]
            if d:
                a["first"] = min(a["first"] or d, d); a["last"] = max(a["last"] or d, d)
        for nm, eff in _record_effects(r):
            info[nm]["net"] += eff
    return info


def sui_preview(db: Database, filename: str, data: bytes, column_overrides: dict | None = None) -> dict:
    parsed = parse_sui(filename, data, column_overrides)
    recs = parsed["records"]
    existing = db.all("SELECT * FROM accounts")
    base = db.get_setting("base_currency", "USD")

    def match_account(name):
        for a in existing:
            names = [a["name"].lower()] + [x.strip().lower() for x in (a.get("aliases") or "").split(",") if x.strip()]
            if name.lower() in names:
                return a["id"]
        return None

    acct_info = _account_usage(recs)
    accounts = []
    for nm, a in sorted(acct_info.items()):
        guessed = guess_account_type_from_usage(nm, a["stats"], a["hint"])
        cur = a["currency"] or ("CNY" if re.search(r"[\u4e00-\u9fff]", nm) else base)
        flags = []
        liab = ledger_is_liability(guessed)
        if not liab and a["net"] < -0.005:
            flags.append("asset ends negative from the file alone - an opening balance is probably missing")
        if liab and a["net"] > 0.005:
            flags.append("liability ends in credit (you'd be owed money) - check the type or the 负债变更 signs")
        accounts.append({"name": nm, "currency": cur, "guessed_type": guessed, "rows": a["rows"],
                         "stats": dict(a["stats"]), "net_from_file": round(a["net"], 2) + 0.0,
                         "first_date": a["first"], "last_date": a["last"], "flags": flags,
                         "existing_account_id": match_account(nm)})
    cats_db = db.all("SELECT * FROM categories")
    by_name = defaultdict(list)
    for c in cats_db:
        by_name[(c["kind"], c["name"].lower())].append(c)
        if c.get("name_zh"):
            by_name[(c["kind"], c["name_zh"].lower())].append(c)
    cats: dict[tuple, dict] = {}
    for r in recs:
        if r.get("type") not in ("expense", "income", "refund"):
            continue
        kind = "income" if r["type"] == "income" else "expense"
        key = (kind, r.get("category") or "", r.get("subcategory") or "")
        if key in cats:
            cats[key]["rows"] += 1; continue
        ex = None
        for c in by_name.get((kind, (r.get("subcategory") or "").lower()), []):
            if c["parent_id"]:
                par = next((p for p in cats_db if p["id"] == c["parent_id"]), None)
                if not r.get("category") or (par and r["category"].lower() in (par["name"].lower(), (par.get("name_zh") or "").lower())):
                    ex = c["id"]; break
        cats[key] = {"kind": kind, "category": key[1], "subcategory": key[2], "rows": 1, "existing_category_id": ex}
    names = lambda f: sorted({str(r[f]) for r in recs if r.get(f)})
    dup = 0
    for r in recs:
        if r.get("date") and r.get("amount") is not None:
            pass
    keys = _dedupe_keys(recs)
    for r, k in zip(recs, keys):
        if k and db.one("SELECT 1 x FROM transactions WHERE external_id=?", (k,)):
            dup += 1
    type_counts = Counter(r["type_raw"] for r in recs)
    missing_fx = sorted({a["currency"] for a in accounts} - set(ledger.fx_rates(db)))
    return {"sheets": parsed["sheets"], "row_count": len(recs), "type_counts": dict(type_counts),
            "unknown_types": sorted({r["type_raw"] for r in recs if not r.get("type")}),
            "accounts": accounts, "categories": sorted(cats.values(), key=lambda c: (c["kind"], c["category"], c["subcategory"])),
            "members": names("member"), "merchants": names("merchant"), "projects": names("project"),
            "duplicates": dup, "missing_fx": missing_fx, "sample": [{k: (str(v) if v is not None else None) for k, v in r.items()} for r in recs[:15]],
            "account_types": {k: v[0] for k, v in __import__("app.db", fromlist=["ACCOUNT_TYPES"]).ACCOUNT_TYPES.items()}}


def _dedupe_keys(recs: list[dict]) -> list[str | None]:
    occ: Counter = Counter()
    out = []
    for r in recs:
        if not r.get("date") or r.get("amount") is None:
            out.append(None); continue
        k = ledger.tx_dedupe_key(r["date"], r["amount"], f"{r.get('account')}>{r.get('account2') or ''}", str(r.get("note") or ""))
        occ[k] += 1
        out.append(f"sui:{k}#{occ[k]}")
    return out


def sui_commit(db: Database, filename: str, data: bytes, options: dict) -> dict:
    """options: account_map {sui_name: {account_id|null, type, currency, name}}, category_map {"kind|cat|sub": id|null},
    skip_duplicates (bool), ledger_id, column_overrides."""
    parsed = parse_sui(filename, data, options.get("column_overrides"))
    recs = parsed["records"]
    account_map: dict = options.get("account_map") or {}
    category_map: dict = options.get("category_map") or {}
    skip_dup = options.get("skip_duplicates", True)
    ledger_id = int(options.get("ledger_id") or 1)
    base = db.get_setting("base_currency", "USD")
    summary = {"imported": 0, "by_type": Counter(), "totals": defaultdict(float), "skipped": [], "duplicates": 0,
               "created_accounts": [], "created_categories": 0}
    acct_ids: dict[str, int] = {}
    created_ids: dict[str, int] = {}
    usage = _account_usage(recs)
    groups = {g["name"]: g["id"] for g in db.all("SELECT * FROM account_groups")}
    from .db import ACCOUNT_TYPES

    def resolve_account(name: str | None, hint: str | None = None) -> int | None:
        if not name:
            return None
        if name in acct_ids:
            return acct_ids[name]
        m = account_map.get(name) or {}
        if m.get("account_id"):
            acct_ids[name] = int(m["account_id"]); return acct_ids[name]
        r = db.one("SELECT id FROM accounts WHERE name=?", (name,))
        if r and not m:
            acct_ids[name] = r["id"]; return r["id"]
        u = usage.get(name)
        t = m.get("type") or (guess_account_type_from_usage(name, u["stats"], hint or u["hint"]) if u else guess_account_type(name, hint))
        if t not in ACCOUNT_TYPES:
            t = "other_asset"
        cur = (m.get("currency") or (u and u["currency"]) or ("CNY" if re.search(r"[\u4e00-\u9fff]", name) else base)).upper()
        with db.tx() as con:
            aid = con.execute("INSERT INTO accounts(group_id,name,type,currency,opening_balance,note) VALUES(?,?,?,?,0,?)",
                              (groups.get(ACCOUNT_TYPES[t][1]), m.get("name") or name, t, cur,
                               "Created by Suishouji import 随手记导入")).lastrowid
        acct_ids[name] = aid
        created_ids[name] = aid
        summary["created_accounts"].append({"name": m.get("name") or name, "type": t, "currency": cur})
        return aid

    cat_cache: dict[str, int | None] = {}

    def resolve_category(kind: str, cat: str | None, sub: str | None) -> int | None:
        key = f"{kind}|{cat or ''}|{sub or ''}"
        if key in cat_cache:
            return cat_cache[key]
        if key in category_map and category_map[key]:
            cat_cache[key] = int(category_map[key]); return cat_cache[key]
        if not cat and not sub:
            cat_cache[key] = None; return None
        with db.tx() as con:
            parent_name = cat or sub
            p = con.execute("SELECT id FROM categories WHERE kind=? AND parent_id IS NULL AND (name=? OR name_zh=?)",
                            (kind, parent_name, parent_name)).fetchone()
            if p:
                pid = p[0]
            else:
                pid = con.execute("INSERT INTO categories(kind,parent_id,name,name_zh,sort) VALUES(?,?,?,?,99)",
                                  (kind, None, parent_name, "")).lastrowid
                summary["created_categories"] += 1
            sub_name = sub or parent_name
            c = con.execute("SELECT id FROM categories WHERE kind=? AND parent_id=? AND (name=? OR name_zh=?)",
                            (kind, pid, sub_name, sub_name)).fetchone()
            if c:
                cid = c[0]
            else:
                cid = con.execute("INSERT INTO categories(kind,parent_id,name,name_zh,sort) VALUES(?,?,?,?,99)",
                                  (kind, pid, sub_name, "")).lastrowid
                summary["created_categories"] += 1
        cat_cache[key] = cid
        return cid

    keys = _dedupe_keys(recs)
    for r, ext in zip(recs, keys):
        where = f"{r['sheet']} row {r['row']}"
        if r.get("error"):
            summary["skipped"].append({"where": where, "reason": r["error"]}); continue
        t = r.get("type")
        if not t:
            summary["skipped"].append({"where": where, "reason": f"unknown type '{r['type_raw']}'"}); continue
        if not r.get("date"):
            summary["skipped"].append({"where": where, "reason": f"bad date '{r.get('date_raw')}'"}); continue
        amt = r.get("amount")
        if amt is None:
            summary["skipped"].append({"where": where, "reason": "no amount"}); continue
        if amt == 0:
            summary["skipped"].append({"where": where, "reason": "zero amount"}); continue
        if not r.get("account"):
            summary["skipped"].append({"where": where, "reason": "no account"}); continue
        if skip_dup and ext and db.one("SELECT 1 x FROM transactions WHERE external_id=?", (ext,)):
            summary["duplicates"] += 1; continue
        roles = LOAN_ROLES.get(t, (None, None))
        hint1 = roles[0] or ("payable" if t == "liability_change" else "receivable" if t == "receivable_change" else None)
        a1 = resolve_account(r["account"], hint1)
        tx: dict[str, Any] = {"date": r["date"], "account_id": a1, "note": str(r.get("note") or ""),
                              "member": r.get("member"), "merchant": r.get("merchant"), "project": r.get("project"),
                              "source": "import", "external_id": ext, "ledger_id": ledger_id, "status": "posted",
                              "raw_text": f"随手记 {r['type_raw']} {r['sheet']}#{r['row']}"}
        if t in ("expense", "income", "refund"):
            kind = "income" if t == "income" else "expense"
            tx["category_id"] = resolve_category(kind, r.get("category"), r.get("subcategory"))
            if t == "expense":
                tx["kind"] = "expense" if amt > 0 else "refund"
            elif t == "income":
                tx["kind"] = "income" if amt > 0 else "expense"
                if amt < 0:
                    tx["category_id"] = None
                    tx["note"] = (tx["note"] + " (negative income in 随手记 export)").strip()
            else:
                tx["kind"] = "refund"
            tx["amount"] = abs(amt)
        elif t in ("transfer",) or t in LOAN_ROLES:
            if not r.get("account2"):
                summary["skipped"].append({"where": where, "reason": f"{r['type_raw']} without second account"}); continue
            tx["kind"] = "transfer"
            tx["to_account_id"] = resolve_account(r["account2"], roles[1])
            tx["amount"] = abs(amt)
            if r.get("to_amount") is not None:
                tx["to_amount"] = abs(r["to_amount"])
            if a1 == tx["to_account_id"]:
                summary["skipped"].append({"where": where, "reason": "transfer to the same account"}); continue
        elif t == "balance":
            tx["kind"], tx["amount"] = "adjustment", amt
        elif t == "receivable_change":
            tx["kind"], tx["amount"] = "adjustment", amt
        elif t == "liability_change":
            tx["kind"], tx["amount"] = "adjustment", -amt  # more debt -> balance goes more negative
        else:
            summary["skipped"].append({"where": where, "reason": f"unsupported type {t}"}); continue
        if not tx["note"]:
            tx["note"] = ""
        try:
            ledger.create_tx(db, tx)
        except ledger.ValidationError as e:
            summary["skipped"].append({"where": where, "reason": str(e)}); continue
        summary["imported"] += 1
        summary["by_type"][r["type_raw"]] += 1
        summary["totals"][r["type_raw"]] += abs(amt)
    # Optional: the balance each account shows in 随手记 today -> set the opening balance so the computed balance matches.
    # For liability accounts the user enters the amount owed (positive), as 随手记 displays it.
    bal_now = ledger.cash_balances(db)
    summary["opening_balances_set"] = []
    for name, aid in created_ids.items():
        target = parse_money((account_map.get(name) or {}).get("balance"))
        if target is None:
            continue
        a = db.one("SELECT type FROM accounts WHERE id=?", (aid,))
        signed = -abs(target) if ledger_is_liability(a["type"]) and target > 0 else target
        opening = round(signed - (bal_now.get(aid, 0.0)), 2)
        # same model as new accounts in the app: an "Opening balance 期初余额" adjustment dated at the account's
        # first day in the file, so today's balance matches 随手记 and history before/after stays consistent
        first = (usage.get(name) or {}).get("first")
        if opening:
            ledger.record_opening_balance(db, aid, opening, first)
        summary["opening_balances_set"].append({"name": name, "opening_balance": opening, "date": first})
    bal_now = ledger.cash_balances(db)
    summary["accounts"] = []
    for name, aid in sorted(acct_ids.items()):
        a = db.one("SELECT name, type, currency, opening_balance FROM accounts WHERE id=?", (aid,))
        ob = next((o["opening_balance"] for o in summary["opening_balances_set"] if o["name"] == name), 0)
        summary["accounts"].append({"sui_name": name, "name": a["name"], "type": a["type"], "currency": a["currency"],
                                    "created": name in created_ids, "opening_balance": (a["opening_balance"] or 0) + ob,
                                    "balance": round(bal_now.get(aid, 0.0), 2) + 0.0})
    have = set(ledger.fx_rates(db))
    summary["missing_fx"] = sorted({a["currency"] for a in summary["accounts"]} - have)
    summary["by_type"] = dict(summary["by_type"])
    summary["totals"] = {k: round(v, 2) for k, v in summary["totals"].items()}
    summary["skipped_count"] = len(summary["skipped"])
    summary["skipped"] = summary["skipped"][:300]
    return summary
